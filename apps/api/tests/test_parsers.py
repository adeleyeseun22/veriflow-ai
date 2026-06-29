from pathlib import Path

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfWriter

from veriflow_api.parsers import ParserError, parse_document, parser_for_extension


def test_parser_registry_supports_ingestion_formats() -> None:
    assert parser_for_extension("pdf").name == "pypdf"
    assert parser_for_extension(".docx").name == "python-docx"
    assert parser_for_extension("xlsx").name == "openpyxl"
    assert parser_for_extension("CSV").name == "python-csv"


def test_parser_registry_rejects_unsupported_extension() -> None:
    with pytest.raises(ParserError, match="No structured parser"):
        parser_for_extension("txt")


def test_csv_parser_extracts_section_and_table(tmp_path: Path) -> None:
    path = tmp_path / "indicators.csv"
    path.write_text("indicator,value\nCoverage,82\nDropout,11\n", encoding="utf-8")

    parsed = parse_document(path, "csv")

    assert parsed.parser_name == "python-csv"
    assert len(parsed.sections) == 1
    assert len(parsed.tables) == 1
    assert parsed.tables[0].column_names == ["indicator", "value"]
    assert parsed.tables[0].rows == [["Coverage", "82"], ["Dropout", "11"]]


def test_docx_parser_preserves_heading_and_table(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    document = WordDocument()
    document.add_heading("Executive Summary", level=1)
    document.add_paragraph("Coverage increased across supported districts.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Indicator"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Coverage"
    table.cell(1, 1).text = "82%"
    document.save(path)

    parsed = parse_document(path, "docx")

    assert parsed.parser_name == "python-docx"
    assert parsed.sections[0].title == "Executive Summary"
    assert "Coverage increased" in parsed.sections[0].content
    assert parsed.tables[0].column_names == ["Indicator", "Value"]
    assert parsed.tables[0].rows == [["Coverage", "82%"]]


def test_xlsx_parser_extracts_each_worksheet(tmp_path: Path) -> None:
    path = tmp_path / "results.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(["State", "Coverage"])
    sheet.append(["FCT", 91])
    second = workbook.create_sheet("Targets")
    second.append(["Indicator", "Target"])
    second.append(["Penta3", 95])
    workbook.save(path)

    parsed = parse_document(path, "xlsx")

    assert parsed.parser_name == "openpyxl"
    assert [section.title for section in parsed.sections] == ["Results", "Targets"]
    assert len(parsed.tables) == 2
    assert parsed.tables[0].rows == [["FCT", "91"]]


def test_pdf_parser_records_pages(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as destination:
        writer.write(destination)

    parsed = parse_document(path, "pdf")

    assert parsed.parser_name == "pypdf"
    assert len(parsed.pages) == 1
    assert parsed.pages[0].page_number == 1
