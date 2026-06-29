from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path

from docx import Document as WordDocument
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from veriflow_api.config import settings
from veriflow_api.parsers.base import ParsedDocument, ParsedSection, ParsedTable, ParserError
from veriflow_api.parsers.utils import (
    normalize_rows,
    normalize_text,
    trim_empty_edges,
    unique_headers,
)

HEADING_PATTERN = re.compile(r"^Heading\s+(\d+)$", re.IGNORECASE)


def iter_blocks(document: DocumentObject) -> Iterator[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


class DocxDocumentParser:
    name = "python-docx"
    version = "1.0"
    extensions = frozenset({"docx"})

    def parse(self, path: Path) -> ParsedDocument:
        try:
            document = WordDocument(str(path))
        except Exception as error:
            raise ParserError("The DOCX file could not be opened for extraction.") from error

        sections: list[ParsedSection] = []
        tables: list[ParsedTable] = []
        heading_path: dict[int, str] = {}
        current_title: str | None = None
        current_level: int | None = None
        current_content: list[str] = []
        current_ordinal: int | None = None

        def flush_section() -> None:
            nonlocal current_content, current_ordinal
            content = normalize_text("\n".join(current_content))
            if not content and current_title is None:
                current_content = []
                return
            ordinal = len(sections)
            path_values = [heading_path[level] for level in sorted(heading_path)]
            sections.append(
                ParsedSection(
                    ordinal=ordinal,
                    title=current_title or "Document body",
                    content=content[: settings.parser_max_section_chars],
                    heading_level=current_level,
                    section_path=path_values,
                    metadata={"section_type": "docx_heading" if current_level else "docx_body"},
                )
            )
            current_ordinal = ordinal
            current_content = []

        for block in iter_blocks(document):
            if isinstance(block, Paragraph):
                text = normalize_text(block.text)
                style_name = block.style.name if block.style is not None else ""
                match = HEADING_PATTERN.match(style_name or "")
                if match and text:
                    flush_section()
                    current_level = max(1, min(int(match.group(1)), 9))
                    current_title = text[:500]
                    heading_path[current_level] = current_title
                    for level in list(heading_path):
                        if level > current_level:
                            heading_path.pop(level)
                    current_ordinal = None
                elif text:
                    current_content.append(text)
            else:
                raw_rows = normalize_rows([[cell.text for cell in row.cells] for row in block.rows])
                rows = trim_empty_edges(raw_rows)
                if not rows:
                    continue
                width = max(len(row) for row in rows)
                headers = unique_headers(rows[0], width)
                data_rows = [row + [""] * (width - len(row)) for row in rows[1:]]
                truncated = len(data_rows) > settings.parser_max_table_rows
                data_rows = data_rows[: settings.parser_max_table_rows]
                section_ordinal = current_ordinal
                if section_ordinal is None and (current_title is not None or current_content):
                    section_ordinal = len(sections)
                tables.append(
                    ParsedTable(
                        ordinal=len(tables),
                        title=f"Table {len(tables) + 1}",
                        source_label="DOCX table",
                        column_names=headers,
                        rows=data_rows,
                        row_count=len(data_rows),
                        column_count=width,
                        is_truncated=truncated,
                        section_ordinal=section_ordinal,
                        metadata={"style": block.style.name if block.style is not None else ""},
                    )
                )

        flush_section()
        properties = document.core_properties
        metadata = {
            "paragraph_count": len(document.paragraphs),
            "source_table_count": len(document.tables),
            "core_properties": {
                "title": properties.title,
                "subject": properties.subject,
                "author": properties.author,
                "keywords": properties.keywords,
            },
        }
        return ParsedDocument(
            parser_name=self.name,
            parser_version=self.version,
            sections=sections,
            tables=tables,
            metadata=metadata,
        )
