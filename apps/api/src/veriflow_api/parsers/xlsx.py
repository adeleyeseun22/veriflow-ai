from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from veriflow_api.config import settings
from veriflow_api.parsers.base import ParsedDocument, ParsedSection, ParsedTable, ParserError
from veriflow_api.parsers.utils import (
    normalize_rows,
    rows_as_text,
    trim_empty_edges,
    unique_headers,
)


class XlsxDocumentParser:
    name = "openpyxl"
    version = "1.0"
    extensions = frozenset({"xlsx"})

    def parse(self, path: Path) -> ParsedDocument:
        try:
            workbook = load_workbook(
                filename=str(path),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except Exception as error:
            raise ParserError("The XLSX workbook could not be opened for extraction.") from error

        sections: list[ParsedSection] = []
        tables: list[ParsedTable] = []
        try:
            for worksheet in workbook.worksheets:
                collected: list[list[str]] = []
                truncated = False
                for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
                    if row_index > settings.parser_max_table_rows:
                        truncated = True
                        break
                    collected.extend(normalize_rows([row]))

                rows = trim_empty_edges(collected)
                if not rows:
                    sections.append(
                        ParsedSection(
                            ordinal=len(sections),
                            title=worksheet.title,
                            content="",
                            section_path=[worksheet.title],
                            metadata={"section_type": "worksheet", "empty": True},
                        )
                    )
                    continue

                width = max(len(row) for row in rows)
                headers = unique_headers(rows[0], width)
                data_rows = [row + [""] * (width - len(row)) for row in rows[1:]]
                section_ordinal = len(sections)
                sections.append(
                    ParsedSection(
                        ordinal=section_ordinal,
                        title=worksheet.title,
                        content=rows_as_text(headers, data_rows),
                        section_path=[worksheet.title],
                        metadata={
                            "section_type": "worksheet",
                            "source_max_row": worksheet.max_row,
                            "source_max_column": worksheet.max_column,
                        },
                    )
                )
                tables.append(
                    ParsedTable(
                        ordinal=len(tables),
                        title=worksheet.title,
                        source_label=f"Worksheet: {worksheet.title}",
                        column_names=headers,
                        rows=data_rows,
                        row_count=len(data_rows),
                        column_count=width,
                        is_truncated=truncated,
                        section_ordinal=section_ordinal,
                        metadata={
                            "worksheet_index": len(sections),
                            "source_max_row": worksheet.max_row,
                            "source_max_column": worksheet.max_column,
                        },
                    )
                )
        finally:
            workbook.close()

        return ParsedDocument(
            parser_name=self.name,
            parser_version=self.version,
            sections=sections,
            tables=tables,
            metadata={"worksheet_count": len(sections)},
        )
